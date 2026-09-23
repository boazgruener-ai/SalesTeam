# -*- coding: utf-8 -*-
"""Builds the files the extension ships for the "Import Research Workbook" dialog, all from sheet_contract.py so they
can never drift from the importer/validator:

  Target_Accounts_Template.xlsx  - empty workbook: the 7 importer sheets + a READ_ME sheet, mandatory columns marked
  import-columns.js              - the mandatory / optional column names per sheet, for the dialog's help text
  research-prompt.txt            - the newest research prompt found in "ChatGPT Market Research" (highest Vn.nn)

    python gen_template.py

Run it after the contract or the prompt changes, then rebuild the package (build_package.py).
"""
import glob
import json
import os
import re

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import sheet_contract as sc

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
PROMPT_DIR = os.path.join(ROOT, "ChatGPT Market Research")

REQUIRED_FILL = PatternFill("solid", fgColor="1F3A5F")
OPTIONAL_FILL = PatternFill("solid", fgColor="E7ECF3")
REQUIRED_FONT = Font(bold=True, color="FFFFFF")
OPTIONAL_FONT = Font(bold=False, color="1F2933")

PRIORITY_VALUES = sc.PRIORITY_BASE + [f"{p} - Provisional" for p in sc.PRIORITY_BASE] + sc.PRIORITY_EXTRA
LISTS = {
    "Companies": {
        "Priority": PRIORITY_VALUES,
        "Evidence_Status": sc.EVIDENCE_STATUS_VALUES,
        "Excluded": ["Yes", "No"],
        "Target_Country_Relationship": ["Local company", "Global company"],
    },
    "Contacts": {"Home_Based": ["Yes", "No"]},
    "Exclusion_List": {"Exclusion_Reason": sc.EXCLUSION_REASONS},
    "Aliases": {"Alias_Type": sc.ALIAS_TYPES},
}


def build_workbook(path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, columns in sc.SHEETS:
        ws = wb.create_sheet(name)
        for i, (header, required, description) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=i, value=header)
            cell.fill = REQUIRED_FILL if required else OPTIONAL_FILL
            cell.font = REQUIRED_FONT if required else OPTIONAL_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            note = ("MANDATORY. " if required else "Optional. ") + (description or "")
            cell.comment = Comment(note.strip(), "SalesTeam")
            ws.column_dimensions[get_column_letter(i)].width = max(14, min(34, len(header) + 4))
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
        headers = [c[0] for c in columns]
        for header, values in LISTS.get(name, {}).items():
            if header not in headers:
                continue
            col = get_column_letter(headers.index(header) + 1)
            dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
            dv.error = "Please pick one of the listed values."
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}2000")

    ws = wb.create_sheet("READ_ME")
    lines = [
        ("SalesTeam - Target Accounts research workbook", True),
        ("", False),
        ("How to use this workbook", True),
        ("1. Fill the sheets by hand, or let an AI assistant fill them using the SalesTeam research prompt (Import Research "
         "Workbook dialog > Copy research prompt).", False),
        ("2. Keep the sheet names and the column headers exactly as they are - the importer finds them by name. "
         "This READ_ME sheet is ignored by the importer.", False),
        ("3. One row per company on Companies. Every other sheet points to its company with Company_ID, which must "
         "exist on Companies.", False),
        ("4. In SalesTeam: Target Accounts menu > Import Research Workbook... > Choose workbook.", False),
        ("", False),
        ("Minimum to import", True),
        ("A Companies sheet with a Company column (one row per account). Add Company_ID to link Contacts, Initiatives, "
         "Investment, Sources and Aliases to their company.", False),
        ("", False),
        ("Mandatory columns of the research format (dark blue headers; hover a header for its description)", True),
    ]
    for name, columns in sc.SHEETS:
        required = [h for h, r, _ in columns if r]
        lines.append((f"{name}: " + ", ".join(required), False))
    lines += [
        ("", False),
        ("Allowed values", True),
        ("Priority: " + ", ".join(PRIORITY_VALUES), False),
        ("Evidence_Status: " + ", ".join(sc.EVIDENCE_STATUS_VALUES), False),
        ("Target_Country_Relationship: Local company, Global company", False),
        ("Exclusion_Reason: " + ", ".join(sc.EXCLUSION_REASONS), False),
        ("Alias_Type: " + ", ".join(sc.ALIAS_TYPES), False),
        ("Dates: YYYY-MM-DD. Numbers: digits only (no thousand separators or currency symbols).", False),
    ]
    for r, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(bold=bold, size=13 if r == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130
    wb.save(path)


def build_columns_js(path):
    data = {
        name: {"mandatory": [h for h, r, _ in cols if r], "optional": [h for h, r, _ in cols if not r]}
        for name, cols in sc.SHEETS
    }
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("// GENERATED by gen_template.py from sheet_contract.py - do not edit by hand.\n")
        f.write("export const IMPORT_COLUMNS = " + json.dumps(data, indent=2, ensure_ascii=False) + ";\n")


def newest_prompt():
    best = None
    for path in glob.glob(os.path.join(PROMPT_DIR, "*.md")):
        m = re.search(r"prompt[_ ]?v(\d+)\.(\d+)", os.path.basename(path), re.I)
        if m and "MASTER" in os.path.basename(path).upper():
            key = (int(m.group(1)), int(m.group(2)))
            if best is None or key > best[0]:
                best = (key, path)
    return best


def copy_prompt(dest):
    best = newest_prompt()
    if not best:
        print("WARNING: no research prompt found in", PROMPT_DIR)
        return None
    text = open(best[1], encoding="utf-8").read()
    with open(dest, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)
    return best


if __name__ == "__main__":
    build_workbook(os.path.join(HERE, "Target_Accounts_Template.xlsx"))
    build_columns_js(os.path.join(HERE, "import-columns.js"))
    best = copy_prompt(os.path.join(HERE, "research-prompt.txt"))
    print("Template, import-columns.js written; prompt:", os.path.basename(best[1]) if best else "NOT FOUND")
