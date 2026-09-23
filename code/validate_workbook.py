# -*- coding: utf-8 -*-
"""Validate a research workbook against the SalesTeam import contract (sheet_contract.py).
Usage: python validate_workbook.py <workbook.xlsx>
Exit code 0 = no ERRORs (WARNINGs allowed), 1 = at least one ERROR.
Read-only: never modifies the workbook.
"""
import sys
import openpyxl
from sheet_contract import (SHEETS, EVIDENCE_STATUS_VALUES, PRIORITY_BASE, PRIORITY_EXTRA,
                            ALIAS_TYPES, EXCLUSION_REASONS, LEGACY_HEADERS)

errors, warnings = [], []


def err(m): errors.append(m)
def warn(m): warnings.append(m)


def cell(v):
    return "" if v is None else str(v).strip()


def read_sheet(wb, name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [cell(h) for h in rows[0]]
    data = [dict(zip(headers, r)) for r in rows[1:] if any(cell(x) for x in r)]
    return headers, data


def priority_ok(p):
    if p in PRIORITY_EXTRA:
        return True
    base = p.split(" - ")[0].strip()
    return base in PRIORITY_BASE


def main(path):
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    names = wb.sheetnames
    data = {}
    for sheet, cols in SHEETS:
        if sheet not in names:
            (err if sheet in ("Companies",) else warn)(f"Sheet '{sheet}' is missing"
                + (" - importer will silently import nothing for it" if sheet != "Companies" else ""))
            continue
        headers, rows = read_sheet(wb, sheet)
        data[sheet] = rows
        legacy_used = [h for h in headers if h in LEGACY_HEADERS.values()]
        if legacy_used:
            warn(f"{sheet}: legacy country-specific header(s) {legacy_used} - accepted, but use the generic names")
        for h, required, _ in cols:
            legacy = LEGACY_HEADERS.get(h)
            if h not in headers and not (legacy and legacy in headers):
                (err if required else warn)(f"{sheet}: column '{h}' missing")
        for h in headers:
            if h and h not in [c[0] for c in cols] and h not in LEGACY_HEADERS.values():
                warn(f"{sheet}: extra column '{h}' (ignored unless the importer knows it)")
        for h, required, _ in cols:
            if required and h in headers:
                empty = sum(1 for r in rows if not cell(r.get(h)))
                if empty:
                    err(f"{sheet}: {empty} row(s) have an empty required '{h}'")

    comp = data.get("Companies", [])
    ids = [cell(r.get("Company_ID")) for r in comp]
    idset = set(ids)
    if len(ids) != len(idset):
        err(f"Companies: {len(ids) - len(idset)} duplicate Company_ID value(s)")
    names_seen = [cell(r.get("Company")).lower() for r in comp]
    if len(names_seen) != len(set(names_seen)):
        warn("Companies: duplicate company names")

    for r in comp:
        cid = cell(r.get("Company_ID"))
        es, pr = cell(r.get("Evidence_Status")), cell(r.get("Priority"))
        if es and es not in EVIDENCE_STATUS_VALUES:
            err(f"Companies {cid}: Evidence_Status '{es}' not in {EVIDENCE_STATUS_VALUES}")
        if pr and not priority_ok(pr):
            err(f"Companies {cid}: Priority '{pr}' not one of {PRIORITY_BASE} (optionally ' - Provisional')")
        if cell(r.get("Excluded")) not in ("Yes", "No"):
            err(f"Companies {cid}: Excluded must be Yes or No, got '{cell(r.get('Excluded'))}'")
        try:
            ec = float(r.get("Evidence_Coverage"))
            if not 0 <= ec <= 1:
                err(f"Companies {cid}: Evidence_Coverage {ec} outside 0-1")
        except (TypeError, ValueError):
            if cell(r.get("Evidence_Coverage")):
                err(f"Companies {cid}: Evidence_Coverage not numeric")
        try:
            ps = float(r.get("Priority_Score"))
            if not 0 <= ps <= 100:
                err(f"Companies {cid}: Priority_Score {ps} outside 0-100")
        except (TypeError, ValueError):
            if cell(r.get("Priority_Score")):
                err(f"Companies {cid}: Priority_Score not numeric")
        rel = cell(r.get("Target_Country_Relationship"))
        if rel and rel not in ("Local company", "Global company"):
            err(f"Companies {cid}: Target_Country_Relationship '{rel}' must be 'Local company' or 'Global company'")
        link = cell(r.get("LinkedIn Link"))
        if link and "linkedin.com/company/" not in link:
            warn(f"Companies {cid}: LinkedIn Link is not a company-page URL: {link}")

    # Company_Type (ownership) must be filled for the active accounts, not only the excluded ones (found in the 2026-09-20 test)
    active = [r for r in comp if cell(r.get("Excluded")).lower() != "yes"]
    untyped = [r for r in active if not cell(r.get("Company_Type"))]
    if active and untyped:
        msg = f"Companies: Company_Type (ownership) empty for {len(untyped)} of {len(active)} active accounts"
        (err if len(untyped) == len(active) else warn)(msg)

    name_by_id = {cell(r.get("Company_ID")): cell(r.get("Company")) for r in comp}
    for sheet in ("Contacts", "Initiatives", "Investment", "Sources", "Exclusion_List", "Aliases"):
        for r in data.get(sheet, []):
            cid = cell(r.get("Company_ID"))
            if cid and cid not in idset:
                err(f"{sheet}: Company_ID '{cid}' does not exist in Companies")
            elif sheet in ("Contacts", "Initiatives", "Investment", "Exclusion_List") and cid:
                nm = cell(r.get("Company"))
                if nm and nm != name_by_id.get(cid):
                    warn(f"{sheet}: Company '{nm}' differs from Companies name '{name_by_id.get(cid)}' for {cid}")
        col = {"Contacts": "Contact_ID", "Initiatives": "Initiative_ID",
               "Investment": "Investment_ID", "Sources": "Source_ID"}.get(sheet)
        if col:
            vals = [cell(r.get(col)) for r in data.get(sheet, [])]
            if len(vals) != len(set(vals)):
                err(f"{sheet}: duplicate {col} values")

    nonstd = {}
    for r in data.get("Aliases", []):
        at = cell(r.get("Alias_Type"))
        if at and at not in ALIAS_TYPES:
            nonstd[at] = nonstd.get(at, 0) + 1
    if nonstd:
        warn(f"Aliases: {len(nonstd)} non-standard Alias_Type value(s) (harmless, free text): "
             + ", ".join(sorted(nonstd)[:8]))
    bad_sp = [r for r in data.get("Aliases", [])
              if cell(r.get("Search_Priority")) not in ("1", "2", "3", "1.0", "2.0", "3.0")]
    if bad_sp:
        err(f"Aliases: {len(bad_sp)} row(s) with Search_Priority not 1, 2 or 3 "
            f"(first: '{cell(bad_sp[0].get('Search_Priority'))}')")
    for r in data.get("Exclusion_List", []):
        er = cell(r.get("Exclusion_Reason"))
        if er and er not in EXCLUSION_REASONS:
            warn(f"Exclusion_List: reason '{er}' not in {EXCLUSION_REASONS}")
    excl_ids = {cell(r.get("Company_ID")) for r in data.get("Exclusion_List", [])}
    flagged = {cell(r.get("Company_ID")) for r in comp if cell(r.get("Excluded")) == "Yes"}
    if excl_ids != flagged:
        warn(f"Excluded=Yes companies ({len(flagged)}) and Exclusion_List rows ({len(excl_ids)}) do not match")

    # coverage summary
    with_c = {cell(r.get("Company_ID")) for r in data.get("Contacts", [])}
    with_i = {cell(r.get("Company_ID")) for r in data.get("Initiatives", [])}
    with_a = {cell(r.get("Company_ID")) for r in data.get("Aliases", [])}
    n = len(comp) or 1
    print(f"Workbook: {path}")
    print("Rows: " + ", ".join(f"{s}={len(data.get(s, []))}" for s, _ in SHEETS))
    print(f"Companies with contacts {len(with_c & idset)}/{len(comp)} ({100*len(with_c & idset)//n}%), "
          f"initiatives {len(with_i & idset)}/{len(comp)}, aliases {len(with_a & idset)}/{len(comp)}")
    for w in warnings[:40]:
        print("WARNING:", w)
    if len(warnings) > 40:
        print(f"... {len(warnings) - 40} more warnings")
    for e in errors[:60]:
        print("ERROR:", e)
    if len(errors) > 60:
        print(f"... {len(errors) - 60} more errors")
    print(f"RESULT: {'FAIL' if errors else 'PASS'} ({len(errors)} errors, {len(warnings)} warnings)")
    return 1 if errors else 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1]))
