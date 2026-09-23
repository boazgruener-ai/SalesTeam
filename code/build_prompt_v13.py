# -*- coding: utf-8 -*-
# Builds Generic_Interactive_B2B_Target_Account_Research_Prompt_V1.3 (.txt + .md + .docx) from V1.2
# plus the workbook contract in sheet_contract.py. V1.2 is never modified.
import os
import re
import docx
from sheet_contract import (SHEETS, EVIDENCE_STATUS_VALUES, PRIORITY_BASE, ALIAS_TYPES,
                            EXCLUSION_REASONS)

D = "C:/Users/41793/OneDrive/Documents/Personal/Claude Code/SalesTeam/ChatGPT Market Research/"
SRC = D + "Generic_Interactive_B2B_Target_Account_Research_Prompt_V1.2.txt"
OUT = D + "Generic_Interactive_B2B_Target_Account_Research_Prompt_V1.3"
BAR = "=" * 70

t = open(SRC, encoding="utf-8").read()


def sub1(old, new):
    global t
    assert t.count(old) == 1, (old, t.count(old))
    t = t.replace(old, new)


sub1("PROMPT — V1.2", "PROMPT — V1.3")

# Step 4: add a pilot size
sub1("Offer:\n100 / 200 / 300 / 400 / 500\n",
     "Offer:\n20 (pilot / test run) / 100 / 200 / 300 / 400 / 500\n\n"
     "A pilot of 10-20 accounts is meant for testing the setup and the workbook format before a full run. "
     "Apply every rule in this prompt to it unchanged; only the universe is smaller.\n")

# Step 1: three opening questions, seller identity
sub1("Ask these ONE AT A TIME.\n\nQuestion 1\n",
     "Ask these ONE AT A TIME.\n\n"
     "Question A — PROJECT NAME\n"
     "“What would you like to call this research project?”\n"
     "Offer a short suggestion, for example the seller's name plus the target market, that the user can accept.\n\n"
     "Question B — SELLER COMPANY\n"
     "“What is the name of the company that will use this research (the seller)?”\n\n"
     "Question C — SELLER COMPANY WEBSITE\n"
     "“What is the company's website address?”\n"
     "Then apply the RESEARCH-FIRST RULE (Section 1.7) immediately, before Question 1.\n\n"
     "Store Project Name, Seller Company and Seller Website in Project_Config. The seller company itself must never be "
     "added to the target-account universe; its subsidiaries/affiliates need explicit user confirmation.\n\n"
     "Question 1\n")

# Section 1.7 research-first rule
sub1("1.6 VISIBLE CONFIGURATION PROGRESS",
     "1.7 RESEARCH-FIRST RULE (PROPOSE, THEN CONFIRM)\n"
     "Once the seller company name and website are known, research them BEFORE asking the remaining setup questions: "
     "read the seller's website (products/services, solutions, industries served, customer references and case studies, "
     "locations, partners, competitors named, company size, leadership) and other credible public sources.\n"
     "Then, for every later setup question where public evidence can suggest an answer, do NOT ask an open question. "
     "Instead show a PROPOSED ANSWER with a one-line basis (for example “from your website's Solutions page”) and ask "
     "the user to Confirm or Adjust. Typical candidates: offering, customer problem, buyer personas, attractive-account "
     "characteristics, geography and primary target country, languages, organization types, industries, company-size "
     "range, likely competitors/partners for the exclusion list, and contact personas.\n"
     "Rules:\n"
     "- Label each proposal with its source (website page, registry, press) and say plainly when it is an inference.\n"
     "- If the evidence is thin or conflicting, say so, propose the most plausible answer anyway, and ask.\n"
     "- The user's answer always overrides the proposal; never re-ask a confirmed item.\n"
     "- Never invent facts about the seller. If the website cannot be read, say so and fall back to asking the question.\n"
     "- Keep the interview short: the goal is that most steps are answered with a single Confirm.\n\n"
     "1.6 VISIBLE CONFIGURATION PROGRESS")

# Section 43 first message
sub1("What products or services are we selling?\nPlease describe the main offering in one or two sentences.”\n\nDo NOT ask all Step 1",
     "What would you like to call this research project?”\n\n"
     "(Then ask for the seller company name and website, run the research-first step, and continue with the offering "
     "question as a proposed answer to confirm.)\n\nDo NOT ask all Step 1")

# Section 30: exact names for the export sheets
sub1("Equivalent names are acceptable if the same state is preserved.",
     "Equivalent names are acceptable for the internal/working sheets if the same state is preserved.\n"
     "EXCEPTION: the seven import sheets defined in Section 45 (Companies, Contacts, Initiatives, Investment, "
     "Sources, Exclusion_List, Aliases) MUST use exactly the sheet names and column headers given there, "
     "because downstream software reads them by name.")

# Section 34: SMOKE TEST command
sub1("RESET / START OVER — begin a new project without destroying prior state\n",
     "RESET / START OVER — begin a new project without destroying prior state\n"
     "SMOKE TEST — see Section 46\n"
     "EXPORT CHECK — verify the current workbook against Section 45 and report every deviation (see Section 46)\n")


def col_lines(cols):
    out = []
    for h, req, desc in cols:
        tag = "REQUIRED" if req else "recommended"
        out.append(f"- {h} [{tag}]" + (f" — {desc}" if desc else ""))
    return "\n".join(out)


sec45 = f"""
{BAR}
45. WORKBOOK OUTPUT CONTRACT (IMPORT SHEETS)
{BAR}

The final workbook is imported by software that reads specific sheets and columns BY EXACT NAME. A sheet or column
with a different name is silently ignored. Therefore, in every delivered workbook version (including checkpoints),
include the seven sheets below in addition to any other sheets from Section 30.

GENERAL RULES
- Sheet names and column headers are in English, exactly as written below (case-sensitive; underscores and the
  spaces in "Alternative Company Name" and "LinkedIn Link" are part of the name), regardless of the language of
  the research, the company names or the conversation. Cell CONTENT (names, titles, descriptions) stays in the
  original language of the source.
- The header row is row 1 of each sheet. One record per row. No merged cells, no title rows above the header, no
  blank rows inside the data, no formulas in these seven sheets (store values only).
- Keep all listed columns even when empty, in the order given. Extra columns may be added AFTER the listed ones.
- Company_ID is unique per company, stable across versions, and is the only link between sheets. Every
  Company_ID used in Contacts, Initiatives, Investment, Aliases and Exclusion_List MUST exist in Companies.
- Never invent contacts, e-mail addresses, figures, dates or URLs. Leave a cell empty when unknown. An empty cell is
  always better than a guess. Dates use YYYY-MM-DD.
- [REQUIRED] columns must be filled in every row of the sheet. [recommended] columns should be filled whenever
  evidence exists.
- "Home" and "Registry" (generic terms, no country names anywhere in the workbook): "Home" means the project's
  PRIMARY TARGET COUNTRY (the "home market" of this research), as confirmed in Step 3 (Geography). Columns starting
  with "Home_" (and Main_Home_Location, Investment_Home, Contacts.Home_Based) hold the figure for the target
  account's operations in that primary target country, as opposed to its Global_* figures. This is independent of
  where the user's own company is headquartered. "Registry" means the official company registry of that primary
  target country (for example Companies House, Registre du Commerce, Handelsregister); Registry_* columns hold that
  registry's data. NEVER put a country or registry name into a sheet name or column header, and never rename these
  headers to a country-specific form, whatever the project's country is. Country names appear only inside cell
  content. If the project has several target countries, use the primary one for Home_*/Registry_* and note the
  others in Research_Status or Priority_Rationale.

SHEET "Companies" — one row per account in the universe (excluded accounts included, with Excluded = Yes)
{col_lines(SHEETS[0][1])}

SHEET "Contacts" — one row per verified person
{col_lines(SHEETS[1][1])}

SHEET "Initiatives" — one row per relevant initiative
{col_lines(SHEETS[2][1])}

SHEET "Investment" — one row per investment/budget evidence item
{col_lines(SHEETS[3][1])}

SHEET "Sources" — one row per source used
{col_lines(SHEETS[4][1])}

SHEET "Exclusion_List" — one row per excluded account (must match Companies.Excluded = Yes)
{col_lines(SHEETS[5][1])}

SHEET "Aliases" — every name variant used for search and matching, including the canonical name itself
{col_lines(SHEETS[6][1])}

ALLOWED VALUES
- Evidence_Status: {" | ".join(EVIDENCE_STATUS_VALUES)}
  Use the Evidence Model of Step 13 to decide the label. Evidence_Coverage is the underlying decimal 0-1.
- Priority: {" | ".join(PRIORITY_BASE)}. Add the suffix " - Provisional" when the evidence is still provisional.
  Use "Insufficient Evidence" (with no Priority_Score) when an account cannot yet be scored responsibly.
- Excluded: Yes | No
- Exclusion_Reason: {" | ".join(EXCLUSION_REASONS)}
- Alias_Type (preferred): {" | ".join(ALIAS_TYPES)}
- Search_Priority: 1 | 2 | 3
- Research_Quality: A | B | C | D
- Contacts.Home_Based: Yes | No (based in the primary target country)

SELF-CHECK BEFORE EVERY DELIVERY
Before delivering a workbook, verify and report: the seven sheet names are exact; no required cell is empty; every
Company_ID reference resolves; Excluded = Yes count equals the Exclusion_List rows; Evidence_Status and Priority use
only the allowed values; no duplicate IDs. Fix problems first; if something cannot be fixed, say so plainly.

{BAR}
46. SMOKE TEST AND EXPORT CHECK
{BAR}

SMOKE TEST
When the user types SMOKE TEST (at any time, with or without a configured project):
- Do NOT do any web research. Do not use or modify any real project data.
- Produce a workbook with ONLY the seven sheets of Section 45, filled with SYNTHETIC data: 3 fictional companies
  (clearly fictional names such as "Example Alpha Ltd"), of which one has Excluded = Yes (Competitor) and one has
  Evidence_Status = Insufficient Evidence with no Priority_Score; 4 fictional contacts across two companies, 2
  initiatives, 1 investment row, 3 sources, 1 exclusion row and 8 or more aliases covering at least four Alias_Type values.
- Fictional contacts must use obviously fake data (e.g. "Test Person One", no real e-mails or URLs; use
  "https://www.linkedin.com/company/example-alpha" style placeholders only for company pages).
- Follow every General Rule and Allowed Value of Section 45 exactly and run the Self-Check.
- Name the file "SMOKE_TEST_workbook.xlsx" and say plainly that it contains synthetic data and must never be
  imported into a real project.

EXPORT CHECK
When the user types EXPORT CHECK: inspect the latest workbook, list every deviation from Section 45 (missing sheet,
misspelled or missing column, empty required cell, unresolved Company_ID, disallowed value), fix them in a new
workbook version, and report what was changed.
"""

t = t.rstrip("\n") + "\n" + sec45

open(OUT + ".txt", "w", encoding="utf-8").write(t)
md = "# Generic Interactive B2B Target Account Research Prompt — V1.3\n\n" + t
open(OUT + ".md", "w", encoding="utf-8").write(md)

doc = docx.Document()
doc.add_heading("Generic Interactive B2B Target Account Research Prompt — V1.3", level=1)
for line in t.split("\n"):
    if re.fullmatch(r"=+", line):
        continue
    doc.add_paragraph(line)
try:
    doc.save(OUT + ".docx")
except PermissionError:
    print("WARNING: .docx is open in Word - close it and re-run to refresh the Word copy")
print("Wrote V1.3 .txt/.md/.docx,", len(t.split("\n")), "lines")
